"""Report what protection a workbook carries, and why a given cell is locked.

Excel blocks editing only when BOTH are true: the sheet is protected, and the
cell itself is locked. Cells are locked by default, so a protected sheet blocks
everything unless specific cells were unlocked first.

Usage:
    .venv/bin/python check_protection.py <workbook.xlsx> [sheet] [cell ...]

Examples:
    .venv/bin/python check_protection.py Audit.xlsx
    .venv/bin/python check_protection.py Audit.xlsx Log G2 G3 G500
"""
import re
import sys
import zipfile

import openpyxl


def report_protection(path):
    with zipfile.ZipFile(path) as zf:
        book = zf.read("xl/workbook.xml").decode("utf-8")
        names = re.findall(r'<sheet\b[^>]*?name="([^"]+)"', book)

        lock = re.search(r"<workbookProtection[^>]*>", book)
        locked_structure = bool(lock and 'lockStructure="1"' in lock.group(0))
        print("Workbook structure protection:", "ON" if locked_structure else "off")
        print("  (the app never changes this)\n")

        print("Sheet protection:")
        parts = sorted(n for n in zf.namelist() if n.startswith("xl/worksheets/sheet"))
        for name, part in zip(names, parts):
            xml = zf.read(part).decode("utf-8")
            prot = re.search(r"<sheetProtection[^>]*/>", xml)
            ranges = xml.count("<protectedRange ")
            extra = f"  (+{ranges} editable range(s))" if ranges else ""
            print(f"  {name:<24} {'PROTECTED' if prot else 'unprotected'}{extra}")


def report_cells(path, sheet, cells):
    wb = openpyxl.load_workbook(path)
    if sheet not in wb.sheetnames:
        print(f"\nNo sheet named {sheet!r}. Available: {wb.sheetnames}")
        wb.close()
        return
    ws = wb[sheet]
    print(f"\nCells on {sheet!r} (sheet protection is "
          f"{'ON' if ws.protection.sheet else 'off'}):")
    for ref in cells:
        cell = ws[ref]
        locked = cell.protection.locked
        blocked = bool(ws.protection.sheet and locked)
        print(f"  {ref:<8} locked={str(locked):<5} value={cell.value!r:<22} "
              f"-> {'BLOCKED by Excel' if blocked else 'editable'}")
    wb.close()


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    target = sys.argv[1]
    print(f"{target}\n")
    report_protection(target)
    if len(sys.argv) > 3:
        report_cells(target, sys.argv[2], sys.argv[3:])
    else:
        print("\nTip: pass a sheet and some cells to see why they're locked, e.g.")
        print("     python check_protection.py file.xlsx Log G2 G3 G500")
